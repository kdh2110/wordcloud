"""Compact app: workbook inspection, isolated jobs and HTML bridge. UI is in index.html."""
# Read-only workbook inspection before expensive NLP.
from io import BytesIO
from zipfile import ZipFile,BadZipFile
from collections import Counter
import openpyxl

MAX_ROWS=50000

def workbook_sheets(content):
    if len(content)>50*1024*1024:raise ValueError('파일 크기는 50MB 이하여야 합니다.')
    try:
        with ZipFile(BytesIO(content)) as z:
            if 'xl/workbook.xml' not in z.namelist():raise ValueError('올바른 Excel(.xlsx) 파일이 아닙니다.')
            if sum(f.file_size for f in z.infolist())>400*1024*1024:raise ValueError('압축 해제한 Excel 크기가 너무 큽니다. 분석할 시트만 새 파일로 저장해 주세요.')
        wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
        try:return wb.sheetnames
        finally:wb.close()
    except (BadZipFile,KeyError) as e:raise ValueError('Excel 파일을 읽을 수 없습니다. .xlsx 형식으로 다시 저장해 주세요.') from e

def inspect_sheet(content,sheet):
    wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
    try:
        rows=wb[sheet].iter_rows(values_only=True)
        first=next(rows,None)
        if first is None:raise ValueError('선택한 시트가 비어 있습니다.')
        header=[str(v).strip() if v is not None else '' for v in first]
        duplicate=[k for k,n in Counter(header).items() if k and n>1]
        if duplicate:raise ValueError('중복된 열 이름이 있습니다: '+', '.join(duplicate))
        if not any(header):raise ValueError('첫 번째 행에 열 이름이 필요합니다.')
        count=0;preview=[]
        for index,row in enumerate(rows,2):
            if index>MAX_ROWS+1:raise ValueError(f'한 시트는 최대 {MAX_ROWS:,}행까지 지원합니다. 불필요한 빈 행을 제거해 주세요.')
            if not any(v is not None for v in row):continue
            count+=1
            if len(preview)<5:preview.append({h:v for h,v in zip(header,row) if h})
        if not count:raise ValueError('분석할 논문 행이 없습니다.')
        return {'columns':[h for h in header if h],'count':count,'preview':preview}
    finally:wb.close()

# Session-owned job folders and subprocess lifecycle; no shared NLP globals.

from dataclasses import dataclass
from pathlib import Path
import atexit,copy,json,os,shutil,subprocess,sys,tempfile,threading,time,uuid

@dataclass
class Job:
    token:str
    owner:str
    root:Path
    process:subprocess.Popen
    config:dict
    started:float
    touched:float
    cancelled:bool=False

    def log_tail(self,limit=16000):
        try:
            with (self.root/'run.log').open('rb') as f:
                f.seek(0,2);size=f.tell();f.seek(max(0,size-limit))
                return f.read().decode('utf-8',errors='replace').replace('\r','\n')
        except FileNotFoundError:return ''

    @property
    def status(self):
        if self.process.poll() is None:return 'running'
        if self.cancelled:return 'cancelled'
        return 'done' if self.process.returncode==0 and (self.root/'success.txt').exists() else 'failed'

    @property
    def output(self):return self.root/'output'/self.config['mode']

class JobManager:
    def __init__(self):
        self.root=Path(tempfile.mkdtemp(prefix='paper_wordcloud_web_'))
        self.jobs={};self.lock=threading.RLock();self.closed=threading.Event()
        self.max_jobs=max(1,int(os.environ.get('WORDCLOUD_MAX_CONCURRENT','1')))
        self.ttl=max(300,int(os.environ.get('WORDCLOUD_SESSION_TTL','86400')))
        self.timeout=max(60,int(os.environ.get('WORDCLOUD_JOB_TIMEOUT','3600')))
        atexit.register(self.close)
        threading.Thread(target=self._reaper,daemon=True).start()

    def start(self,owner,content,config):
        with self.lock:
            if any(j.owner==owner and j.status=='running' for j in self.jobs.values()):
                raise ValueError('현재 진행 중인 분석을 완료하거나 중지한 뒤 다시 실행해 주세요.')
            if sum(j.status=='running' for j in self.jobs.values())>=self.max_jobs:
                raise ValueError('서버에서 다른 분석을 진행 중입니다. 잠시 후 다시 시작해 주세요.')
            token=uuid.uuid4().hex;root=self.root/owner/token;root.mkdir(parents=True)
            (root/'input.xlsx').write_bytes(content)
            cfg=copy.deepcopy(config);cfg['job_dir']=str(root);cfg['cache_dir']=str(self.root/owner/'cache')
            env=os.environ.copy();env['PYTHONUNBUFFERED']='1';env['PYTHONIOENCODING']='utf-8'
            for name in ['OMP_NUM_THREADS','OPENBLAS_NUM_THREADS','MKL_NUM_THREADS']:
                env.setdefault(name,'1')
            logfile=(root/'run.log').open('wb')
            try:
                process=subprocess.Popen([sys.executable,'-u',str(Path(__file__).resolve()),'--worker'],
                    stdin=subprocess.PIPE,stdout=logfile,stderr=subprocess.STDOUT,env=env,
                    cwd=str(Path(__file__).parent))
                process.stdin.write(json.dumps(cfg,ensure_ascii=False).encode('utf-8'));process.stdin.close()
            except Exception:
                if 'process' in locals():process.terminate();process.wait(timeout=10)
                shutil.rmtree(root,ignore_errors=True)
                raise
            finally:logfile.close()
            now=time.time();job=Job(token,owner,root,process,cfg,now,now);self.jobs[token]=job
            return job

    def get(self,owner,token):
        with self.lock:
            job=self.jobs.get(token)
            if job is None or job.owner!=owner:return None
            job.touched=time.time();return job

    def cancel(self,owner,token):
        with self.lock:
            j=self.get(owner,token)
            if j and j.status=='running':
                j.cancelled=True;j.process.terminate()
                try:j.process.wait(timeout=5)
                except subprocess.TimeoutExpired:j.process.kill();j.process.wait(timeout=5)

    def clear(self,owner):
        with self.lock:
            for token,j in list(self.jobs.items()):
                if j.owner==owner:
                    self.cancel(owner,token);del self.jobs[token]
            shutil.rmtree(self.root/owner,ignore_errors=True)

    def _reaper(self):
        while not self.closed.wait(30):
            with self.lock:
                now=time.time()
                for j in list(self.jobs.values()):
                    if j.status=='running' and now-j.started>self.timeout:
                        self.cancel(j.owner,j.token)
                        with (j.root/'run.log').open('a',encoding='utf-8') as f:f.write('\nERROR|최대 실행 시간을 초과하여 분석을 중지했습니다.\n')
                active={j.owner for j in self.jobs.values() if j.status=='running'}
                for owner in {j.owner for j in self.jobs.values()}-active:
                    if all(now-j.touched>self.ttl for j in self.jobs.values() if j.owner==owner):self.clear(owner)

    def close(self):
        if self.closed.is_set():return
        self.closed.set()
        with self.lock:
            for owner in {j.owner for j in self.jobs.values()}:self.clear(owner)
            shutil.rmtree(self.root,ignore_errors=True)

# One isolated Python process per analysis: model/JVM and globals never mix.
from pathlib import Path
import sys,json,traceback,zipfile,os

def memory_snapshot():
    parts=[]
    try:
        for line in Path('/proc/self/status').read_text().splitlines():
            if line.startswith('VmRSS:'):
                parts.append(f"worker_rss_mb={int(line.split()[1])/1024:.1f}")
        for filename,label in [('memory.current','container_mb'),('memory.max','container_limit_mb')]:
            path=Path('/sys/fs/cgroup')/filename
            if path.exists():
                value=path.read_text().strip()
                parts.append(f"{label}={int(value)/1048576:.1f}" if value.isdigit() else f"{label}={value}")
    except (OSError,ValueError):
        pass
    return ' '.join(parts) or 'memory_metrics_unavailable'


def start_memory_monitor():
    stop=threading.Event()
    def monitor():
        while not stop.is_set():
            print('MEM|'+memory_snapshot(),flush=True)
            stop.wait(15)
    threading.Thread(target=monitor,daemon=True).start()
    return stop


def worker_main():
    cfg=json.loads(sys.stdin.read())
    root=Path(cfg['job_dir']).resolve()
    print('STEP|0|실행 환경과 입력 데이터를 확인하고 있습니다.',flush=True)
    import pipeline as p
    print('STEP|0|입력 Excel을 읽고 필수 열과 피인용 값을 확인합니다.',flush=True)
    # Config values are data only; users cannot choose code, output paths or fonts.
    mapping=cfg['columns']
    for key,value in mapping.items():
        if key in {'TITLE_COL','ABSTRACT_COL','KEYWORD_COL','CITATION_COL','AUTHORS_COL','YEAR_COL','JOURNAL_COL','DOI_COL','LINK_COL','DOCUMENT_ID_COL','SOURCE_COL'}:
            setattr(p,key,value)
    p.SHEET_NAME=cfg['sheet']
    for key in ['TITLE_WEIGHT','ABSTRACT_WEIGHT','KEYWORD_WEIGHT','CITATION_ALPHA','CITATION_AS_OF','WORDCLOUD_MAX_WORDS','WORDCLOUD_MIN_DF','MIN_PHRASE_FREQ','MIN_PHRASE_DF','PMI_THRESHOLD','MAX_NGRAM','WORDCLOUD_COLORMAP','WORDCLOUD_WIDTH','WORDCLOUD_HEIGHT','WORDCLOUD_MAX_FONT','WORDCLOUD_MIN_FONT','WORDCLOUD_SIZE_POWER','HTML_INCLUDE_ABSTRACT','KEYWORD_SEPARATOR']:
        if key in cfg['settings']:setattr(p,key,cfg['settings'][key])
    if not p.CITATION_COL:
        p.CITATION_ALPHA=0.0;p.RANKING_VERSION='no_idf'
    p.EXCLUDE_GENERIC_PHRASES.update(cfg.get('exclude',[]))
    p.MISSING_CITATIONS='neutral'
    font=os.environ.get('WORDCLOUD_FONT_PATH','').strip()
    if font:p.WORDCLOUD_FONT_PATH=font
    if min(p.TITLE_WEIGHT,p.ABSTRACT_WEIGHT,p.KEYWORD_WEIGHT)<0 or sum([p.TITLE_WEIGHT,p.ABSTRACT_WEIGHT,p.KEYWORD_WEIGHT])<=0:
        raise ValueError('가중치는 0 이상이며, 한 가지 이상은 0보다 커야 합니다.')
    if p.WORDCLOUD_MIN_FONT>p.WORDCLOUD_MAX_FONT:raise ValueError('최소 글자 크기는 최대 글자 크기보다 작아야 합니다.')
    # Preflight before loading spaCy; invalid citation values retain Excel row numbers.
    header,records=p.read_rows(root/'input.xlsx',cfg['mode'])
    needs_ko=any(p.HANGUL.search(p.clean(r['title'])+' '+p.clean(r['abstract'])+' '+p.clean(r['keywords'])) for r in records)
    if needs_ko:
        import jpype
        try:jpype.getDefaultJVMPath()
        except Exception as e:raise RuntimeError('한국어 분석에는 Java(JDK/JRE)가 필요합니다. 설치 후 앱을 다시 시작하거나 Docker 실행 방법을 이용하세요.') from e
    import importlib.util
    if importlib.util.find_spec('en_core_web_sm') is None:
        raise RuntimeError('영어 모델이 없습니다. python -m spacy download en_core_web_sm 을 실행해 주세요.')
    # Printing stages as they are reached. No fabricated completion percentage.
    print('STEP|1|저자키워드 사전과 제목·초록을 분석합니다.',flush=True)
    original_report=p.write_df_report
    def reporting(*a,**kw):
        print('STEP|3|점수 계산, Excel 및 네 가지 워드클라우드를 생성합니다.',flush=True)
        return original_report(*a,**kw)
    p.write_df_report=reporting
    out=p.run(root/'input.xlsx',cfg['mode'],root/'output',Path(cfg['cache_dir']),preloaded=(header,records))
    paths=[out/'analysis.xlsx',out/'paper_explorer.html']+[out/f'wordcloud_{v}.png' for v in p.SCORE_FIELDS]
    if not (out/'analysis.xlsx').exists() or not (out/'paper_explorer.html').exists():
        raise RuntimeError('필수 결과 파일이 생성되지 않았습니다. 실행 기록을 확인해 주세요.')
    for path in paths:
        if not path.exists():print(f'WARNING|{path.name} 생성 실패. 나머지 결과는 다운로드할 수 있습니다.',flush=True)
    (root/'success.txt').write_text(str(out.relative_to(root)),encoding='utf-8')
    print('STEP|4|분석과 결과 파일 생성이 완료되었습니다.',flush=True)


# Display an existing v8 HTML report, adding a download for its selected PNG.
from pathlib import Path
import json,re

# This script runs AFTER the existing report script, so D/version/changeVersion
# refer to the original report state. It does not recompute any rankings.
DOWNLOAD_SCRIPT = r'''
<script>
(function () {
  "use strict";
  const button = document.getElementById("downloadCloud");
  function updateCloudDownload() {
    const cloud = D.clouds[version];
    button.disabled = !(cloud && typeof cloud.png === "string" && cloud.png.startsWith("data:image/png;base64,"));
    button.title = button.disabled ? "현재 조건의 이미지가 없습니다." : "현재 인용·IDF 조건의 원본 PNG 저장";
  }
  button.addEventListener("click", function () {
    const cloud = D.clouds[version];
    if (!cloud || !cloud.png || !cloud.png.startsWith("data:image/png;base64,")) return;
    const a = document.createElement("a");
    a.href = cloud.png;
    a.download = "wordcloud_" + version + ".png";
    document.body.appendChild(a);
    a.click();
    a.remove();
  });
  const originalChangeVersion = changeVersion;
  changeVersion = function (v) {
    originalChangeVersion(v);
    updateCloudDownload();
  };
  const csv = document.getElementById("csv");
  if (csv) {
    csv.textContent = "리스트 다운로드";
    csv.title = "선택한 표현의 검색·발견 위치·정렬 조건에 맞는 전체 논문 목록을 CSV로 저장합니다. 현재 페이지뿐 아니라 모든 페이지를 포함합니다.";
  }
  updateCloudDownload();
})();
</script>
'''

def prepare_report(text):
    """Accept only a trusted administrator-provided full v8 report, not a template."""
    found=re.search(r'<script\b(?=[^>]*\bid=[\"\']dataset[\"\'])[^>]*>(.*?)</script\s*>',text,re.S|re.I)
    if not found:
        raise ValueError('report.html에 분석 데이터가 없습니다. v8에서 생성한 paper_explorer.html을 넣어 주세요.')
    try:data=json.loads(found.group(1))
    except (ValueError,TypeError) as e:
        raise ValueError('HTML의 분석 데이터를 읽을 수 없습니다. __PAYLOAD__가 있는 템플릿 대신 완성된 결과 HTML을 사용하세요.') from e
    if not isinstance(data,dict) or not {'papers','terms','clouds'}.issubset(data):
        raise ValueError('v8 결과 HTML 형식이 아닙니다.')
    if not isinstance(data['papers'],list) or not isinstance(data['terms'],list) or not isinstance(data['clouds'],dict):
        raise ValueError('논문·표현·이미지 데이터 형식을 확인해 주세요.')
    if not re.search(r'function\s+changeVersion\s*\(',text):
        raise ValueError('인용·IDF 전환 함수가 없습니다. v8 결과 HTML을 사용하세요.')
    # Always apply the current styles, even to an already-patched report.
    style_pattern = r'<style\b[^>]*id=["\']viewer-ui-v9["\'][^>]*>.*?</style\s*>'
    if re.search(style_pattern, text, flags=re.S|re.I):
        text = re.sub(style_pattern, lambda m: COMPACT_STYLE.strip(), text, flags=re.S|re.I)
    else:
        text = re.sub(r'</head\s*>', lambda m: COMPACT_STYLE.strip() + m.group(), text, count=1, flags=re.I)
    if re.search(r'id=[\"\']downloadCloud[\"\']',text):
        return text
    marker=re.search(r'<p\b(?=[^>]*\bid=[\"\']cloudnote[\"\'])[^>]*>.*?</p\s*>',text,re.S|re.I)
    if not marker:raise ValueError('워드클라우드 영역을 찾을 수 없습니다. HTML에 cloudnote 요소가 필요합니다.')
    control='\n<div class="cloud-download"><button id="downloadCloud" type="button">워드클라우드 다운로드</button></div>\n'
    text=text[:marker.end()]+control+text[marker.end():]
    css='\n<style>.cloud-download{display:flex;justify-content:flex-end;margin:4px 0 14px}.cloud-download button{font-size:12px;padding:6px 10px}@media print{.cloud-download{display:none}}</style>\n'
    text=re.sub(r'</head\s*>',lambda m:css+m.group(),text,count=1,flags=re.I)
    text=re.sub(r'</body\s*>',lambda m:DOWNLOAD_SCRIPT+m.group(),text,count=1,flags=re.I)
    return text

# Scoped inside the HTML iframe. Size values also override older v8 CSS.
COMPACT_STYLE = r'''
<style id="viewer-ui-v9">
#paperSort,#csv,#downloadCloud {
  font-family:inherit!important;font-size:12px!important;line-height:1.4!important;
  padding:5px 8px!important;height:30px!important;min-height:30px!important;
  width:auto!important;max-width:100%;box-sizing:border-box!important;
  margin:0!important;white-space:nowrap!important;flex:0 0 auto!important;
}
#paperSort {min-width:96px!important;}
#csv,#downloadCloud {min-width:0!important;}
.list-actions {display:flex;align-items:center;justify-content:flex-end;gap:6px;flex-wrap:wrap;}
.cloud-download {display:flex;justify-content:flex-end;margin:4px 0 14px;}
@media print {.cloud-download{display:none}}
</style>
'''


# ---- Native HTML UI / server adapter ----
import base64,math,hashlib

ALIASES={
'TITLE_COL':['Title','제목','논문제목','Article Title'],'ABSTRACT_COL':['Abstract','초록'],
'KEYWORD_COL':['Author Keywords','저자키워드','저자 키워드'],
'CITATION_COL':['Cited by','Citation Count','Citations','Times Cited, WoS Core','피인용횟수','피인용 횟수'],
'AUTHORS_COL':['Authors','저자','Author Full Names'],'YEAR_COL':['Year','Publication Year','연도'],
'JOURNAL_COL':['Source title','Publication Title','저널명','학술지명'],'DOI_COL':['DOI'],
'LINK_COL':['Link','URL'],'DOCUMENT_ID_COL':['EID','UT (Unique WOS ID)','UT'],'SOURCE_COL':['Source','출처']}


def detect_keyword_separator(content,sheet,column):
    """Prefer explicit list separators; commas only when no stronger delimiter exists."""
    wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
    try:
        rows=wb[sheet].iter_rows(values_only=True)
        headers=[str(v).strip() if v is not None else '' for v in next(rows)]
        index=headers.index(column);counts={';':0,'|':0,'\n':0,',':0}
        for n,row in enumerate(rows):
            if n>=300:break
            value=str(row[index] or '') if index<len(row) else ''
            value=re.sub(r'\([^)]*\)|\[[^]]*\]','',value)
            for sep in counts:counts[sep]+=sep in value
        for sep in [';','|','\n']:
            if counts[sep]:return sep
        return ',' if counts[','] else ';'
    finally:wb.close()


def inspect_upload(content,sheet=None):
    sheets=workbook_sheets(content)
    # Scan headers only, then inspect the first matching sheet's rows once.
    wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
    matches=[]
    try:
        for name in sheets:
            if sheet is not None and name!=sheet:continue
            row=next(wb[name].iter_rows(values_only=True),())
            columns=[str(v).strip() for v in row if v is not None]
            mapping={k:next((c for c in columns if c.casefold() in {a.casefold() for a in aliases}),None) for k,aliases in ALIASES.items()}
            if all(mapping[k] for k in ['TITLE_COL','ABSTRACT_COL','KEYWORD_COL']):matches.append((name,mapping))
    finally:wb.close()
    if not matches:raise ValueError('제목·초록·저자키워드 열을 찾지 못했습니다. 첫 행의 열 이름을 Title/Abstract/Author Keywords 또는 제목/초록/저자키워드로 맞춰 주세요.')
    sheet,mapping=matches[0];info=inspect_sheet(content,sheet)
    separator=detect_keyword_separator(content,sheet,mapping['KEYWORD_COL'])
    notice=f'여러 데이터 시트 중 첫 번째 「{sheet}」를 사용합니다.' if len(matches)>1 else ''
    if mapping['CITATION_COL'] is None:notice+=' 피인용 횟수 열이 없어 인용 가중치는 적용하지 않습니다.'
    return dict(kind='inspected',sheet=sheet,columns=info['columns'],count=info['count'],mapping=mapping,separator=separator,notice=notice.strip())


def validated_config(req,content,filename,info):
    # Removed UI options are server-owned fixed settings, not client parameters.
    columns=dict(info['mapping'])
    bounds={'TITLE_WEIGHT':(0,100,5),'ABSTRACT_WEIGHT':(0,100,1),'KEYWORD_WEIGHT':(0,100,5),'CITATION_ALPHA':(0,10,.5)}
    settings={};values=req.get('settings',{})
    for k,(lo,hi,default) in bounds.items():
        v=float(values.get(k,default))
        if not math.isfinite(v) or not lo<=v<=hi:raise ValueError(f'{k} 설정은 {lo}~{hi} 범위여야 합니다.')
        settings[k]=v
    if sum(settings[k] for k in ['TITLE_WEIGHT','ABSTRACT_WEIGHT','KEYWORD_WEIGHT'])<=0:raise ValueError('한 가지 이상의 가중치는 0보다 커야 합니다.')
    color=values.get('WORDCLOUD_COLORMAP','Dark2')
    if color not in {'hsv','rainbow','tab10','Dark2','viridis'}:raise ValueError('색상 설정을 확인해 주세요.')
    settings.update(WORDCLOUD_COLORMAP=color,KEYWORD_SEPARATOR=info['separator'],CITATION_AS_OF='',MAX_NGRAM=5,HTML_INCLUDE_ABSTRACT=True,
        WORDCLOUD_MAX_WORDS=50,WORDCLOUD_WIDTH=1200,WORDCLOUD_HEIGHT=800,MIN_PHRASE_FREQ=5,MIN_PHRASE_DF=3,PMI_THRESHOLD=3.0,WORDCLOUD_MIN_DF=2)
    if columns['CITATION_COL'] is None:settings['CITATION_ALPHA']=0.0
    return dict(sheet=info['sheet'],mode='full',source_name=filename,source_hash=hashlib.sha256(content).hexdigest(),columns=columns,settings=settings,exclude=[v.strip().lower() for v in str(req.get('exclude','')).splitlines() if v.strip()])


def native_report(text,template):
    text=prepare_report(text)
    # Refresh controls in older generated reports as well as plain v8 reports.
    patterns=[r'<style id="wc-native-css">.*?</style>',r'<div id="wc-create-wrap">.*?</div>',
              r'<div id="wc-status".*?</div>',r'<dialog id="wc-modal".*?</dialog>',r'<script id="wc-native-script">.*?</script>']
    css,button,status,modal,script=[re.search(pattern,template,re.S).group() for pattern in patterns]
    for pattern in patterns:text=re.sub(pattern,'',text,flags=re.S)
    if not re.search(r'</header\s*>',text,re.I):raise ValueError('결과 HTML의 header 영역이 없습니다.')
    text=re.sub(r'</head\s*>',lambda m:css+m.group(),text,count=1,flags=re.I)
    text=re.sub(r'</header\s*>',lambda m:button+m.group()+status,text,count=1,flags=re.I)
    return re.sub(r'</body\s*>',lambda m:modal+script+m.group(),text,count=1,flags=re.I)


def serve():
    import streamlit as st
    import streamlit.components.v1 as components
    st.set_page_config(page_title='연구 키프레이즈 · 논문 탐색',page_icon='📚',layout='wide',initial_sidebar_state='collapsed')
    st.markdown('<style>[data-testid="stMainBlockContainer"],.block-container{padding-top:4.5rem!important;padding-bottom:1rem;max-width:1900px}</style>',unsafe_allow_html=True)
    @st.cache_resource
    def get_manager():return JobManager()
    mgr=get_manager();ss=st.session_state
    if 'owner' not in ss:ss.owner=uuid.uuid4().hex
    owner=ss.owner
    component=components.declare_component('paper_report_native',path=str(Path(__file__).with_name('ui')))
    # Events are consumed once; only bytes from this session can be analyzed.
    req=ss.get('next_request')
    if isinstance(req,dict) and req.get('id')!=ss.get('handled_request'):
        ss.handled_request=req.get('id');action=req.get('action');ss.response={'kind':'idle'}
        try:
            latest=mgr.get(owner,ss.get('job_token'))
            if latest and latest.status=='running' and action not in {'poll','fetch','cancel'}:raise ValueError('진행 중인 분석을 완료하거나 중지해 주세요.')
            if action=='upload':
                encoded=req.get('data','')
                if not isinstance(encoded,str) or len(encoded)>70_000_000:raise ValueError('파일 크기는 50MB 이하여야 합니다.')
                content=base64.b64decode(encoded,validate=True)
                info=inspect_upload(content)
                ss.upload_bytes=content;ss.upload_filename=str(req.get('filename','input.xlsx'))[:255];ss.sheet_info=info;ss.response=info
            elif action=='sheet':
                if 'upload_bytes' not in ss:raise ValueError('파일을 먼저 선택해 주세요.')
                info=inspect_upload(ss.upload_bytes,req.get('sheet'));ss.sheet_info=info;ss.response=info
            elif action=='analyze':
                if 'upload_bytes' not in ss or 'sheet_info' not in ss:raise ValueError('파일을 먼저 선택해 주세요.')
                config=validated_config(req,ss.upload_bytes,ss.upload_filename,ss.sheet_info)
                job=mgr.start(owner,ss.upload_bytes,config);ss.job_token=job.token
                ss.pop('upload_bytes',None);ss.pop('sheet_info',None)
            elif action=='cancel':mgr.cancel(owner,ss.get('job_token'))
            elif action=='clear':
                mgr.clear(owner)
                for key in ['job_token','active_job','handled_job','upload_bytes','sheet_info']:ss.pop(key,None)
                ss.response={'kind':'idle','message':'기본 결과로 돌아왔습니다.'}
            elif action=='excel':
                active=mgr.get(owner,ss.get('active_job'))
                if active: ss.response={'kind':'download','download':base64.b64encode((active.output/'analysis.xlsx').read_bytes()).decode()}
            elif action not in {'poll','fetch'}:raise ValueError('알 수 없는 요청입니다.')
        except Exception as e:ss.response={'kind':'error','message':str(e)}
        ss.pop('next_request',None)
        if isinstance(req,dict): req.pop('data',None)
    job=mgr.get(owner,ss.get('job_token'))
    response=dict(ss.get('response',{'kind':'idle'}))
    if job and job.status=='running':
        log=job.log_tail();steps=re.findall(r'STEP\|\d+\|([^\n]+)',log);elapsed=int(time.time()-job.started)
        response.update(kind='running',message=(steps[-1] if steps else '분석 준비 중')+f' · {elapsed//60}분 {elapsed%60}초')
        # Surface worker telemetry in Community Cloud logs during UI polling.
        bucket=(job.token,elapsed//15)
        if ss.get('memory_log_bucket')!=bucket:
            ss.memory_log_bucket=bucket
            metrics=re.findall(r'MEM\|([^\n]+)',log)
            print('ANALYSIS|'+response['message']+(' | '+metrics[-1] if metrics else ''),flush=True)
    elif job and ss.get('handled_job')!=job.token:
        ss.handled_job=job.token
        if job.status=='done':
            ss.active_job=job.token;response.update(kind='done',message='새 결과로 바뀌었습니다.')
        elif job.status=='cancelled':response.update(kind='cancelled',message='분석을 중지했습니다. 기존 결과를 유지합니다.')
        else:
            log=job.log_tail()
            errors=re.findall(r'ERROR\|([^\n]+)',log)
            code=job.process.returncode
            hint=('프로세스가 강제 종료되었습니다. 메모리 부족 또는 외부 종료 가능성이 있습니다.'
                  if code in {-9,137} else '분석 프로세스가 중단되었습니다.')
            message=errors[-1] if errors else f'{hint} 종료 코드: {code}. Manage app의 로그를 확인해 주세요.'
            print(f'WORKER_EXIT|code={code}\n{log}',flush=True)
            response.update(kind='error',message=message)
        ss.response=response
    active=mgr.get(owner,ss.get('active_job'));response['hasOwn']=bool(active);response['request_id']=ss.get('handled_request','initial')
    if active and response['kind'] not in {'running','error','cancelled'}:response['message']='내 분석 결과 · '+active.config['source_name']+' · '+('최대 100건 샘플' if active.config['mode']=='sample' else '전체 논문')
    path=active.output/'paper_explorer.html' if active else Path(__file__).with_name('report.html')
    if not active and not path.exists():path=Path(__file__).with_name('sample_report.html')
    try:
        stamp=(str(path),path.stat().st_mtime_ns,(Path(__file__).with_name('ui')/'index.html').stat().st_mtime_ns)
        if ss.get('html_stamp')!=stamp:
            template=re.search(r'<!--BEGIN_REPORT-->\n(.*?)\n<!--END_REPORT-->',(Path(__file__).with_name('ui')/'index.html').read_text(encoding='utf-8'),re.S).group(1)
            ss.report_html=native_report(path.read_text(encoding='utf-8-sig'),template);ss.report_key=hashlib.sha256(ss.report_html.encode()).hexdigest();ss.html_stamp=stamp
        send_html=ss.get('sent_key')!=ss.report_key or (isinstance(req,dict) and req.get('action')=='fetch')
        event=component(report_html=ss.report_html if send_html else '',report_key=ss.report_key,state=response,height=1300,key='native_report',default=None)
        ss.sent_key=ss.report_key
        if isinstance(event,dict) and event.get('id')!=ss.get('handled_request'):
            ss.next_request=event;st.rerun()
    except (OSError,ValueError,AttributeError) as e:st.error(f'결과 HTML 확인이 필요합니다: {e}')

if __name__=='__main__':
    if '--worker' in sys.argv:
        monitor_stop=start_memory_monitor()
        try:worker_main()
        except Exception as e:
            print(f'ERROR|{type(e).__name__}: {e}',flush=True);traceback.print_exc();sys.exit(1)
        finally:monitor_stop.set()
    else:serve()