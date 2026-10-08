"""v8: citation-weighted clickable paper explorer + 필드별 DF 점수 + 저자키워드 + 겹침 검사 대각선 워드클라우드."""
from __future__ import annotations
import argparse, collections, gzip, hashlib, html, importlib.metadata
import json, math, re, sqlite3, time, unicodedata
import tempfile, base64, io, gc
from datetime import datetime
from urllib.parse import urlparse
from contextlib import contextmanager, nullcontext
from pathlib import Path
from functools import lru_cache

# ── 조정 가능한 설정 ──────────────────────────────────────────────
TITLE_COL = 'Title'
ABSTRACT_COL = 'Abstract'
KEYWORD_COL = 'Author Keywords'
SHEET_NAME = 0
AUTHORS_COL = 'Authors'
YEAR_COL = 'Year'
JOURNAL_COL = 'Source title'
DOI_COL = 'DOI'
LINK_COL = 'Link'
DOCUMENT_ID_COL = 'EID'
SOURCE_COL = 'Source'
HTML_INCLUDE_ABSTRACT = True
KEYWORD_SEPARATOR = ';'       # 첨부 파일에서 확인: / 또는 쉼표로 분리하지 않음
SAMPLE_SIZE = 100
RANDOM_SEED = 42
MODEL = 'en_core_web_sm'
BATCH_SIZE = 8
CHUNK_SIZE = 16                # 본문: 논문 수 / 저자키워드: 표현 수
MIN_WORD_FREQ = 1              # 최종 결과 표시 기준. 후보 발굴 기준과 별개
MIN_PHRASE_FREQ = 5
MIN_PHRASE_DF = 3
PMI_THRESHOLD = 3.0            # 모든 이분할의 PMI 중 최솟값(log2)
MAX_NGRAM = 5
INCLUDE_VERBS = False
INCLUDE_ADJECTIVES = False     # False여도 복합어 안의 형용사는 유지
MIN_KO_LENGTH = 2
KEEP_KO_ORTHOGRAPHIC_COMPOUNDS = True
REMOVE_COPYRIGHT_TAIL = True
EXTRA_REMOVE_TEXT = ['양식의 맨 위', '양식의 맨 아래']
EXTRA_STOPWORDS = set()        # 소문자 영어 lemma 또는 한국어 기본형
REJECT_PHRASES = {'case report', 'systematic review', 'meta-analysis',
                  'validation study', 'evaluation study'}
# 자동 후보를 검토한 뒤 추가 가능. N-gram/표제어 정규화된 표현으로 입력
REJECT_AUTO_PHRASES = set()
ACCEPT_PHRASES = []            # 저자키워드와 같은 방식으로 보호할 검토 완료 용어
MAX_EXCEL_ROWS = 1_048_575      # 헤더를 제외한 Excel 한 시트의 최대 행 수
VERSION = '8.0.0'
MORPH_CACHE_VERSION = '1.0.0'
NORMALIZATION_VERSION = 'ai-parenthesis-apostrophe-v2'
STOP_EN = set('''study research article paper work result method approach analysis
finding conclusion use provide show demonstrate investigate evaluate assess examine
identify develop propose report significant significantly difference effect author
background objective purpose introduction discussion copyright reserve right'''.split())
STOP_KO = {'연구', '논문', '결과', '방법', '분석', '목적', '결론'}
HANGUL = re.compile('[가-힣]')
KO_RUN = re.compile(r'[가-힣]+(?:[ \t]+[가-힣]+)*')
NOUNS = {'NOUN', 'PROPN'}
BAD_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean(value):
    if value is None:
        return ''
    s = html.unescape(str(value))
    s = re.sub(r'<[^>]+>', '', s)  # B<sub>12</sub> -> B12
    s = unicodedata.normalize('NFKC', s)
    s = s.translate(str.maketrans({"’":"'", "‘":"'", "ʼ":"'"}))
    s = s.translate(str.maketrans({'‐':'-', '‑':'-', '–':'-', '−':'-'}))
    if s.strip().lower() in {'[no abstract available]', 'nan', 'none', 'n/a'}:
        return ''
    if REMOVE_COPYRIGHT_TAIL:
        s = s.split('©', 1)[0]
    for x in EXTRA_REMOVE_TEXT:
        s = s.replace(x, ' ')
    s = re.sub(r'\bartificial\s+intelligence\s*\(\s*a\.?\s*i\.?\s*\)',
               'artificial intelligence', s, flags=re.IGNORECASE)
    return re.sub(r'\s+', ' ', BAD_XML.sub('', s)).strip()


def language(s):
    ko, en = bool(HANGUL.search(s)), bool(re.search('[A-Za-z]', s))
    return 'MIXED' if ko and en else 'KO' if ko else 'EN' if en else 'OTHER'


def tok(text, lemma, pos, lang, start, end, stop=False, parts=None):
    return dict(text=text, lemma=lemma, pos=pos, lang=lang, start=start,
                end=end, stop=bool(stop), parts=parts or [lemma])


class Cache:
    """원문 정제 결과+분석기 버전별 SQLite 캐시. 오류는 캐시하지 않는다."""
    def __init__(self, path, signature):
        self.db = sqlite3.connect(str(path))
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('CREATE TABLE IF NOT EXISTS cache (k TEXT PRIMARY KEY, v TEXT)')
        self.signature = signature

    def key(self, s):
        return hashlib.sha256((self.signature+'\0'+s).encode()).hexdigest()

    def get(self, s):
        r = self.db.execute('SELECT v FROM cache WHERE k=?', (self.key(s),)).fetchone()
        return json.loads(r[0]) if r else None

    def put(self, s, v):
        self.db.execute('INSERT OR REPLACE INTO cache VALUES (?,?)',
                        (self.key(s), json.dumps(v, ensure_ascii=False)))

    def commit(self):
        self.db.commit()

    def close(self):
        self.db.commit()
        self.db.close()


class Analyzer:
    def __init__(self, cache_dir, needs_ko):
        import spacy
        try:
            self.nlp = spacy.load(MODEL, exclude=['parser', 'ner'])
        except OSError as e:
            raise RuntimeError('영어 모델 필요: python -m spacy download en_core_web_sm') from e
        # tagger / attribute_ruler / lemmatizer / tok2vec를 유지한다.
        self.nlp.max_length = max(self.nlp.max_length, 2_000_000)
        self.okt = None
        if needs_ko:
            from konlpy.tag import Okt
            self.okt = Okt(max_heap_size=512)
        versions = {p: importlib.metadata.version(p) for p in ['spacy']}
        try:
            versions['konlpy'] = importlib.metadata.version('konlpy')
        except importlib.metadata.PackageNotFoundError:
            versions['konlpy'] = None
        signature = json.dumps([MORPH_CACHE_VERSION, MODEL, self.nlp.meta.get('version'), versions,
                                KEEP_KO_ORTHOGRAPHIC_COMPOUNDS], sort_keys=True)
        self.cache = Cache(Path(cache_dir)/'morphology.sqlite', signature)
        self.cache_hits = 0

    @lru_cache(maxsize=512)
    def ko_pos(self, s):
        # norm=False, stem=False: 원문 위치를 보존하여 거짓 인접 결합 방지
        return tuple(self.okt.pos(s, norm=False, stem=False))

    def ko_tokens(self, s, offset):
        out, cursor = [], 0
        for form, tag in self.ko_pos(s):
            start = s.find(form, cursor)
            if start < 0:
                raise ValueError(f'한국어 형태소 원문 위치 정렬 실패: {form!r}')
            end = start+len(form)
            pos = {'Noun':'NOUN','Verb':'VERB','Adjective':'ADJ'}.get(tag, tag)
            lemma = form
            if tag in {'Verb','Adjective'}:
                stemmed = self.okt.pos(form, norm=False, stem=True)
                if len(stemmed) == 1:
                    lemma = stemmed[0][0]
            t = tok(form, lemma, pos, 'KO', offset+start, offset+end)
            # 붙여 쓴 연속 명사만 복원: '인공/지능' -> '인공지능'. 조사 경계는 넘지 않음.
            if (KEEP_KO_ORTHOGRAPHIC_COMPOUNDS and out and pos == 'NOUN'
                    and out[-1]['pos'] == 'NOUN' and out[-1]['end'] == t['start']):
                old = out[-1]
                old['text'] += form
                old['lemma'] += lemma
                old['end'] = t['end']
                old['parts'].append(lemma)
            else:
                out.append(t)
            cursor = end
        return out

    def en_tokens(self, doc, offset):
        out = []
        for t in doc:
            if t.is_space:
                continue
            lemma = (t.lemma_ or t.text).casefold()
            out.append(tok(t.text, lemma, t.pos_, 'EN', offset+t.idx,
                           offset+t.idx+len(t.text), t.is_stop))
        # 'all-cause' 등 붙어 있는 하이픈 연결어는 하나로 보존한다.
        merged, i = [], 0
        while i < len(out):
            j = i
            while (j+2 < len(out) and out[j+1]['text'] == '-'
                   and out[j]['end'] == out[j+1]['start']
                   and out[j+1]['end'] == out[j+2]['start']
                   and re.search('[A-Za-z0-9]', out[j]['text'])
                   and re.search('[A-Za-z0-9]', out[j+2]['text'])):
                j += 2
            if j > i:
                q = out[i:j+1]
                surface = ''.join(x['text'] for x in q)
                lemma = ''.join(x['lemma'] for x in q)
                # 숫자가 섞인 연령/수치 결합어는 자동 명사구 후보로 승격하지 않는다.
                pos = q[-1]['pos'] if not re.search(r'\d', surface) else 'X'
                merged.append(tok(surface, lemma, pos, 'EN', q[0]['start'], q[-1]['end']))
            else:
                merged.append(out[i])
            i = j+1
        return merged

    def many(self, texts):
        # Only Python dict/list/str values escape this scope; no spaCy objects.
        zone = self.nlp.memory_zone() if hasattr(self.nlp, 'memory_zone') else nullcontext()
        with zone:
            return self._many(texts)

    def _many(self, texts):
        """한 chunk만 메모리에 유지. 배치 실패 시 개별 텍스트로 재시도."""
        texts = list(dict.fromkeys(texts))
        results, errors, pending, jobs = {}, {}, {}, []
        for s in texts:
            hit = self.cache.get(s)
            if hit is not None:
                results[s] = hit
                self.cache_hits += 1
                continue
            pending[s] = []
            try:
                cursor = 0
                for match in KO_RUN.finditer(s):
                    if match.start() > cursor:
                        jobs.append((s[cursor:match.start()], (s, cursor)))
                    pending[s].extend(self.ko_tokens(match.group(), match.start()))
                    cursor = match.end()
                if cursor < len(s):
                    jobs.append((s[cursor:], (s, cursor)))
            except Exception as e:
                errors[s] = f'{type(e).__name__}: {e}'
        # 배치 오류 시 개별 재시도로 정상 논문의 처리를 계속한다.
        for start in range(0, len(jobs), BATCH_SIZE):
            batch = [(t, ctx) for t, ctx in jobs[start:start+BATCH_SIZE] if ctx[0] not in errors]
            try:
                docs = list(self.nlp.pipe(batch, as_tuples=True, batch_size=BATCH_SIZE, n_process=1))
            except Exception:
                docs = []
                for t, ctx in batch:
                    try:
                        docs.append((self.nlp(t), ctx))
                    except Exception as e:
                        errors[ctx[0]] = f'{type(e).__name__}: {e}'
            for doc, (s, offset) in docs:
                try:
                    pending[s].extend(self.en_tokens(doc, offset))
                except Exception as e:
                    errors[s] = f'{type(e).__name__}: {e}'
        for s, ts in pending.items():
            if s not in errors:
                ts.sort(key=lambda t:t['start'])
                self.cache.put(s, ts)
                results[s] = ts
        self.cache.commit()
        return results, errors


def word_key(t):
    return t['lang'], t['lemma']


def phrase_name(key):
    # 한국어 띄어쓰기 용어는 원래 토큰 경계(공백)를 보존한다.
    return ' '.join(x[1] for x in key)


def content_token(t):
    stops = STOP_KO if t['lang'] == 'KO' else STOP_EN
    return (not t['stop'] and t['lemma'] not in stops | EXTRA_STOPWORDS
            and bool(re.search('[A-Za-z가-힣]', t['lemma'])))


def candidate_token(t):
    return content_token(t) and t['pos'] in NOUNS | ({'ADJ'} if t['lang']=='EN' else set())


def standalone(t):
    allowed = NOUNS | ({'VERB'} if INCLUDE_VERBS else set()) | ({'ADJ'} if INCLUDE_ADJECTIVES else set())
    return (content_token(t) and t['pos'] in allowed
            and len(t['lemma']) >= (MIN_KO_LENGTH if t['lang']=='KO' else 2))


def contiguous(a, b, source):
    return (a['lang'] == b['lang'] and
            (source[a['end']:b['start']] == '' or source[a['end']:b['start']].isspace()))


def windows(ts, source, n):
    for i in range(len(ts)-n+1):
        w = ts[i:i+n]
        if all(candidate_token(t) for t in w) and all(contiguous(a,b,source) for a,b in zip(w,w[1:])):
            yield tuple(word_key(t) for t in w), w


class Trie:
    def __init__(self):
        self.root = {}

    def add(self, key):
        node = self.root
        for part in key:
            node = node.setdefault(part, {})
        node[None] = key

    def longest(self, ts, i, source):
        node, best, j = self.root, None, i
        while j < len(ts) and word_key(ts[j]) in node:
            # 실제 중간 토큰은 생략하지 않는다. 공백 이외 누락문자도 넘지 않는다.
            if j > i and not (source[ts[j-1]['end']:ts[j]['start']] == '' or
                               source[ts[j-1]['end']:ts[j]['start']].isspace()):
                break
            node = node[word_key(ts[j])]
            j += 1
            if None in node:
                best = (node[None], j)
        return best


def make_seeds(raw_keywords, analyzer, errors):
    from tqdm.auto import tqdm
    seeds, labels = set(), collections.defaultdict(set)
    values = sorted(set(raw_keywords) | {clean(x) for x in ACCEPT_PHRASES})
    for start in tqdm(range(0,len(values),CHUNK_SIZE), desc='1/4 저자키워드 사전'):
        batch = values[start:start+CHUNK_SIZE]
        results, failed = analyzer.many(batch)
        for s, message in failed.items():
            errors.append(dict(excel_row='', field=KEYWORD_COL, text=s, error=message))
        for s, ts in results.items():
            if not ts or language(s) == 'OTHER':
                continue
            key = tuple(word_key(t) for t in ts)
            name = phrase_name(key)
            if name in REJECT_PHRASES or name in STOP_EN | STOP_KO | EXTRA_STOPWORDS:
                continue
            seeds.add(key)
            labels[key].add(s)
    return seeds, labels


# 피인용 횟수는 원본 Excel의 수집 시점 값. 외부 API를 조회하지 않는다.
CITATION_COL = 'Cited by'
CITATION_ALPHA = 0.5      # 0: 인용 미반영 / 0.2: 약하게 / 0.5: 기본 / 1: 강하게
MISSING_CITATIONS = 'neutral'  # neutral: 가중치 1 + 경고 / error: 누락 시 중단
CITATION_AS_OF = ''      # 원본 데이터를 수집한 날짜(예: '2026-09-30'), 모르면 공란
SCORE_FIELDS = {'no_idf':'base', 'idf':'with_idf',
                'citation':'citation_score', 'citation_idf':'citation_with_idf'}


def parse_citations(value, excel_row):
    """빈값과 실제 0을 구분한다. 음수/비정수/잘못된 문자열은 조용히 0으로 바꾸지 않는다."""
    missing = value is None or (isinstance(value,str) and value.strip().lower() in
                               {'', 'n/a', 'na', 'nan', 'none', 'null', '-'})
    if isinstance(value,float) and math.isnan(value): missing=True
    if missing:
        if MISSING_CITATIONS=='error':
            raise ValueError(f'Excel {excel_row}행: {CITATION_COL} 값이 비어 있습니다.')
        return None
    if isinstance(value,bool):
        raise ValueError(f'Excel {excel_row}행: 피인용 횟수에 True/False를 사용할 수 없습니다.')
    text=str(value).strip()
    if ',' in text:
        if not re.fullmatch(r'\d{1,3}(?:,\d{3})+(?:\.0+)?',text):
            raise ValueError(f'Excel {excel_row}행: 잘못된 피인용 횟수 {value!r}')
        text=text.replace(',','')
    from decimal import Decimal, InvalidOperation
    try: number=Decimal(text)
    except InvalidOperation:
        raise ValueError(f'Excel {excel_row}행: 피인용 횟수는 0 이상의 정수여야 합니다: {value!r}') from None
    if not number.is_finite() or number<0 or number!=number.to_integral_value() or number>2**53-1:
        raise ValueError(f'Excel {excel_row}행: 피인용 횟수는 0 이상의 정수여야 합니다: {value!r}')
    return int(number)


class PaperHits:
    """Store every term/document match on disk; read only one term at a time."""
    def __init__(self, path):
        self.db = sqlite3.connect(str(path))
        self.db.execute('PRAGMA cache_size=-4096')
        self.db.execute('PRAGMA temp_store=FILE')
        self.db.execute('CREATE TABLE hits (lang TEXT, term TEXT, row INTEGER, mask INTEGER, surfaces TEXT, PRIMARY KEY(lang,term,row)) WITHOUT ROWID')
        self.pending = 0

    def add_document(self, row, matches):
        self.db.executemany('INSERT INTO hits VALUES (?,?,?,?,?)',
            ((lang, term, row, h[0], json.dumps(h[1], ensure_ascii=False))
             for (lang, term), h in matches.items()))
        self.pending += 1
        if self.pending >= 64:
            self.commit()

    def commit(self):
        self.db.commit()
        self.pending = 0

    def get(self, key, default=None):
        values = {row: [mask, json.loads(surfaces)] for row, mask, surfaces in
            self.db.execute('SELECT row,mask,surfaces FROM hits WHERE lang=? AND term=? ORDER BY row', key)}
        return values if values else default

    def items(self):
        from itertools import groupby
        cursor = self.db.execute('SELECT lang,term,row,mask,surfaces FROM hits ORDER BY lang,term,row')
        try:
            for key, group in groupby(cursor, key=lambda x: (x[0], x[1])):
                yield key, {row: [mask, json.loads(surfaces)]
                            for _, _, row, mask, surfaces in group}
        finally:
            cursor.close()

    def close(self):
        self.db.close()


def citation_term_stats(context):
    """집계된 용어의 논문별 필드 존재 여부에 인용 가중치를 적용한다."""
    by_row={r['excel_row']:r for r in context['records']}
    result={}
    for key,matches in context['hits'].items():
        contributions=[]; bases=[]; total=0; cited=0; missing=0
        for row,h in matches.items():
            mask=h[0]
            base=(TITLE_WEIGHT if mask&1 else 0)+(ABSTRACT_WEIGHT if mask&2 else 0)+(KEYWORD_WEIGHT if mask&4 else 0)
            c=by_row[row]['citations']
            missing += c is None
            if c is not None:
                total += c
                cited += c>0
            bases.append(base)
            contributions.append(base*(1+CITATION_ALPHA*math.log1p(c or 0)))
        result[key]=dict(base_check=math.fsum(bases),citation_score=math.fsum(contributions),
                         citation_total=total,cited_papers=cited,citation_missing=missing)
    return result


def read_rows(path, mode):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[SHEET_NAME] if isinstance(SHEET_NAME,int) else wb[SHEET_NAME]
        rows = ws.iter_rows(values_only=True)
        header = [str(v).strip() if v is not None else '' for v in next(rows)]
        for name in [TITLE_COL,ABSTRACT_COL,KEYWORD_COL]+([CITATION_COL] if CITATION_COL else []):
            if header.count(name) != 1:
                raise ValueError(f'컬럼 {name!r}이 없거나 중복입니다. 실제 컬럼: {header}')
        records = []
        for excel_row, values in enumerate(rows,2):
            data = dict(zip(header, values))
            if not any(v is not None for v in values):
                continue
            records.append(dict(excel_row=excel_row, citations=parse_citations(data.get(CITATION_COL),excel_row),
                 title=str(data[TITLE_COL] or ''), abstract=str(data[ABSTRACT_COL] or ''),
                 keywords=str(data[KEYWORD_COL] or ''), id=str(data.get(DOCUMENT_ID_COL) or data.get(DOI_COL) or excel_row),
                 authors=str(data.get(AUTHORS_COL) or ''), year=str(data.get(YEAR_COL) or ''),
                 journal=str(data.get(JOURNAL_COL) or ''), doi=str(data.get(DOI_COL) or ''),
                 link=str(data.get(LINK_COL) or ''), source=str(data.get(SOURCE_COL) or '')))
    finally:
        wb.close()
    if mode == 'sample' and len(records) > SAMPLE_SIZE:
        import random
        records = sorted(random.Random(RANDOM_SEED).sample(records,SAMPLE_SIZE),key=lambda r:r['excel_row'])
    if not records:
        raise ValueError('분석할 데이터가 없습니다.')
    return header, records


@contextmanager
def record_writer(path):
    conn=sqlite3.connect(path)
    try:
        conn.execute('CREATE TABLE records (id INTEGER PRIMARY KEY, payload TEXT)')
        yield conn
        conn.commit()
    finally:
        conn.close()


def dump_line(conn, obj):
    conn.execute('INSERT INTO records(payload) VALUES (?)',
                 (json.dumps(obj,ensure_ascii=False),))


def iter_records(path):
    conn=sqlite3.connect(path)
    try:
        for (payload,) in conn.execute('SELECT payload FROM records ORDER BY id'):
            yield json.loads(payload)
    finally:
        conn.close()


def discover(path):
    """SQLite로 빈도 누적. 모든 N-gram을 Python 메모리에 보관하지 않는다.
    PMI는 동일 길이 창의 이분할 분포를 사용. 길이 n의 prefix/suffix 주변확률을
    각각 같은 길이 n의 적격 창에서 세므로 분모와 표본공간이 일치한다.
    """
    from tqdm.auto import tqdm
    dbpath = path.parent/'candidate_work.sqlite'
    db = sqlite3.connect(dbpath)
    db.execute('DROP TABLE IF EXISTS grams')
    db.execute('CREATE TABLE grams (k TEXT PRIMARY KEY,n INTEGER,c INTEGER,df INTEGER)')
    for record in tqdm(iter_records(path),desc='3/4 전체 corpus 후보 수집'):
        counts = collections.Counter()
        for field in record['fields'].values():
            if field['error']:
                continue
            for n in range(2,MAX_NGRAM+1):
                for key,w in windows(field['tokens'],field['clean'],n):
                    counts[(json.dumps(key,ensure_ascii=False),n)] += 1
        db.executemany('INSERT INTO grams VALUES (?,?,?,1) ON CONFLICT(k) DO UPDATE SET c=c+excluded.c,df=df+1',
                       [(k,n,c) for (k,n),c in counts.items()])
        db.commit()
    db.execute('DROP TABLE IF EXISTS margins')
    db.execute('CREATE TABLE margins (n INTEGER,split INTEGER,side TEXT,k TEXT,c INTEGER,PRIMARY KEY(n,split,side,k))')
    totals = collections.Counter()
    # 주변 분포도 SQLite에 나누어 저장하여 메모리 증가를 줄인다.
    cursor = db.execute('SELECT k,n,c FROM grams')
    while True:
        batch = cursor.fetchmany(2000)
        if not batch:
            break
        margins = collections.Counter()
        for serial,n,c in batch:
            key = json.loads(serial)
            totals[n] += c
            for split in range(1,n):
                margins[(n,split,'L',json.dumps(key[:split],ensure_ascii=False))] += c
                margins[(n,split,'R',json.dumps(key[split:],ensure_ascii=False))] += c
        db.executemany('INSERT INTO margins VALUES (?,?,?,?,?) ON CONFLICT(n,split,side,k) DO UPDATE SET c=c+excluded.c',
                       [(*k,c) for k,c in margins.items()])
    db.commit()
    candidates, accepted = [], set()
    for serial,n,c,df in db.execute('SELECT k,n,c,df FROM grams WHERE c>=?',(MIN_PHRASE_FREQ,)):
        key = tuple(tuple(x) for x in json.loads(serial))
        pmi_values = []
        for split in range(1,n):
            m = []
            for side,part in [('L',key[:split]),('R',key[split:])]:
                m.append(db.execute('SELECT c FROM margins WHERE n=? AND split=? AND side=? AND k=?',
                     (n,split,side,json.dumps(part,ensure_ascii=False))).fetchone()[0])
            pmi_values.append(math.log2(c*totals[n]/(m[0]*m[1])))
        pmi = min(pmi_values)
        # 끝 명사 조건은 원본 토큰에서 별도로 확인한다(동일 lemma의 품사 변동 가능).
        candidates.append(dict(key=key, freq=c, df=df, pmi=pmi,
             eligible=df>=MIN_PHRASE_DF and pmi>=PMI_THRESHOLD and
                      phrase_name(key) not in REJECT_PHRASES | REJECT_AUTO_PHRASES))
    db.close()
    # 후보별 실제 명사 종결 관찰. 빈도는 품사 적격 인접 창의 겹침을 허용한 빈도다.
    wanted = {x['key'] for x in candidates}
    noun_ended = set()
    for record in iter_records(path):
        for field in record['fields'].values():
            if field['error']:
                continue
            for n in range(2,MAX_NGRAM+1):
                for key,w in windows(field['tokens'],field['clean'],n):
                    if key in wanted and w[-1]['pos'] in NOUNS:
                        noun_ended.add(key)
    for row in candidates:
        row['eligible'] = row['eligible'] and row['key'] in noun_ended
        if row['eligible']:
            accepted.add(row['key'])
    return candidates, accepted


def write_excel(path, tables):
    # 사용자 요청 라이브러리 openpyxl. write_only로 대용량 메모리 증가 억제.
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(write_only=True)
    for name, headers, rows in tables:
        ws, count, page = None,0,0
        for row in rows:
            if ws is None or count >= MAX_EXCEL_ROWS:
                page += 1
                suffix = '' if page == 1 else f'_{page}'
                ws = wb.create_sheet(name[:31-len(suffix)]+suffix)
                ws.freeze_panes = 'A2'
                cells=[]
                for h in headers:
                    cell=WriteOnlyCell(ws,value=h)
                    cell.font=Font(color='FFFFFF',bold=True)
                    cell.fill=PatternFill('solid',fgColor='203864')
                    cells.append(cell)
                ws.append(cells)
                count=0
            cells=[]
            for v in row:
                if isinstance(v,str):
                    v=BAD_XML.sub('',v)
                    if len(v)>32700:
                        v=v[:32600]+' … [Excel 셀 길이 제한으로 생략]'
                    cell=WriteOnlyCell(ws,value=v)
                    cell.data_type='s'  # 원문이 = 로 시작해도 수식으로 해석하지 않음
                    cells.append(cell)
                else:
                    cells.append(v)
            ws.append(cells)
            count+=1
        if ws is None:
            ws=wb.create_sheet(name)
            ws.append(headers)
    wb.save(path)


# ── DF 워드클라우드 설정: 형태소 분석/복합어 발굴 이후에만 적용 ──
REPORT_VERSION = "5.0-integrated-diagonal"
EXCLUDE_EXACT_TERMS = {
    'model', 'patient', 'group', 'datum', 'level', 'time', 'year', 'case',
    'association', 'factor', 'performance', 'system', 'strategy', 'potential',
    'impact', 'condition', 'response', 'need', 'evaluation', 'expression',
    'value', 'number', 'type', 'process',
}
WORDCLOUD_MIN_DF = 2
WORDCLOUD_MAX_WORDS = 50
WORDCLOUD_WIDTH = 1200
WORDCLOUD_HEIGHT = 800
WORDCLOUD_FONT_PATH = None          # None: 한글 폰트 자동 탐색
WORDCLOUD_BACKGROUND = 'white'
WORDCLOUD_COLORMAP = 'Dark2'
WORDCLOUD_RANDOM_STATE = 42


def normalize_term(term):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(term))).strip().casefold()


def choose_wordcloud_font(frequencies):
    if WORDCLOUD_FONT_PATH:
        path = Path(WORDCLOUD_FONT_PATH)
        if not path.is_file():
            raise FileNotFoundError(f'지정 폰트가 없습니다: {path}')
        return str(path)
    candidates = [
        '/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf',
        '/usr/share/fonts/truetype/nanum/NanumGothic.ttf',
        '/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf',
        'C:/Windows/Fonts/malgunbd.ttf',
        'C:/Windows/Fonts/malgun.ttf',
        '/System/Library/Fonts/Supplemental/AppleGothic.ttf',
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
    if any(HANGUL.search(term) for term in frequencies):
        raise RuntimeError('한글 폰트가 필요합니다. Colab 설치 셀을 실행하거나 WORDCLOUD_FONT_PATH를 지정하세요.')
    from matplotlib import font_manager
    return font_manager.findfont(font_manager.FontProperties(family='DejaVu Sans', weight='bold'))


def reprocess_existing(input_path, output_dir='nlp_output_df'):
    """Aggregated workbooks do not retain term-to-paper membership."""
    raise ValueError('클릭형 HTML에는 논문별 연결 정보가 필요합니다. 원본 논문 Excel을 full/sample 모드로 실행하세요.')


def run(input_path, mode='sample', output_dir='nlp_output_df', cache_dir='nlp_cache', preloaded=None):
    # Temporary SQLite holds intermediate records; it is removed even on failure.
    with tempfile.TemporaryDirectory(prefix='paper_nlp_v5_') as work_dir:
        return _run_impl(input_path, mode, output_dir, cache_dir, Path(work_dir), preloaded)


def _run_impl(input_path, mode, output_dir, cache_dir, work_dir, preloaded=None):
    from tqdm.auto import tqdm
    started=time.perf_counter()
    mode=str(mode).strip().lower()
    if mode not in {'sample','full'}:
        raise ValueError('mode는 sample 또는 full')
    out=Path(output_dir)/mode
    out.mkdir(parents=True,exist_ok=True)
    Path(cache_dir).mkdir(parents=True,exist_ok=True)
    if not isinstance(CITATION_ALPHA,(int,float)) or not math.isfinite(CITATION_ALPHA) or CITATION_ALPHA<0:
        raise ValueError('CITATION_ALPHA는 0 이상의 유한수여야 합니다.')
    if MISSING_CITATIONS not in {'neutral','error'}:
        raise ValueError('MISSING_CITATIONS는 neutral 또는 error')
    header,records=preloaded if preloaded is not None else read_rows(input_path,mode)
    missing=sum(r['citations'] is None for r in records)
    print(f'피인용 횟수: {CITATION_COL} / α={CITATION_ALPHA:g} / 누락 {missing}건')
    if missing: print('주의: 피인용 횟수 누락은 실제 0과 구분하여 기록하며, 인용 가중치만 1로 적용합니다.')
    keyword_values=[]
    keyword_rows=collections.defaultdict(list)
    for r in records:
        for value in r['keywords'].split(KEYWORD_SEPARATOR):
            value=clean(value)
            if value:
                keyword_values.append(value)
                keyword_rows[value].append(r['excel_row'])
    needs_ko=any(HANGUL.search(clean(r['title'])+' '+clean(r['abstract'])) for r in records)
    needs_ko=needs_ko or any(HANGUL.search(x) for x in keyword_values+ACCEPT_PHRASES)
    print(f'{len(records):,}건 / mode={mode} / 한국어 분석기 필요={bool(needs_ko)}')
    print('후보 기준:',MIN_PHRASE_FREQ,MIN_PHRASE_DF,PMI_THRESHOLD,'(빈도/DF/PMI)')
    analyzer=Analyzer(cache_dir,needs_ko)
    errors=[]
    paper_hits=None
    cache_hits=analyzer.cache_hits
    model_version=analyzer.nlp.meta.get('version')
    try:
        seeds,labels=make_seeds(keyword_values,analyzer,errors)
        keyword_raw,keyword_df,keyword_by_row=keyword_counters(records,labels)
        for e in errors:
            e['excel_row']=', '.join(map(str, keyword_rows.get(e['text'],[]))) or '사용자 사전'
        del keyword_values, keyword_rows
        # 본문 캐시는 seed 사전 및 PMI 기준과 독립적이다. 전체 실행 때 재사용 가능.
        rawpath=work_dir/'morphology.sqlite'
        field_errors=0
        with record_writer(rawpath) as f:
            for start in tqdm(range(0,len(records),CHUNK_SIZE),desc='2/4 제목·초록 형태소 분석'):
                batch=records[start:start+CHUNK_SIZE]
                texts=[clean(r[k]) for r in batch for k in ['title','abstract']]
                analyzed,failed=analyzer.many(texts)
                for r in batch:
                    fields={}
                    for k in ['title','abstract']:
                        s=clean(r[k]); err=failed.get(s,'')
                        fields[k]=dict(clean=s,tokens=analyzed.get(s,[]),error=err)
                        if err:
                            field_errors+=1
                            errors.append(dict(excel_row=r['excel_row'],field=k,text='',error=err))
                    dump_line(f,dict(**r,fields=fields))
        # Morphology is complete: downstream stages consume SQLite token records.
        cache_hits=analyzer.cache_hits
        analyzer.cache.close()
        analyzer.ko_pos.cache_clear()
        analyzer=None
        del analyzed, failed, texts, fields, batch
        gc.collect()
        print('STEP|2|전체 논문의 자동 후보를 선별하고 빈도·DF를 집계합니다.', flush=True)
        candidates,auto=discover(rawpath)
        seed_trie,auto_trie=Trie(),Trie()
        for key in seeds: seed_trie.add(key)
        for key in auto: auto_trie.add(key)
        frequencies=collections.defaultdict(lambda:collections.Counter())
        dfs=collections.Counter()
        field_dfs = {f: collections.Counter() for f in ('title', 'abstract')}
        union_dfs=collections.Counter()
        seed_hits=collections.Counter()
        used_sources=collections.defaultdict(set)
        orthographic=collections.Counter()
        orthographic_parts={}
        paper_hits=PaperHits(work_dir/'paper_hits.sqlite')
        for rec in tqdm(iter_records(rawpath),total=len(records),desc='4/4 복합어 결합·빈도 계산'):
            doc_terms=set(); counts={}; used=set()
            document_hits={}
            for field,data in rec['fields'].items():
                ts,source=data['tokens'],data['clean']
                final=[]; i=0
                while i<len(ts):
                    # 같은 시작점에서는 검증된 seed 우선, 각 사전 내 최장 일치.
                    hit=seed_trie.longest(ts,i,source)
                    origin='저자키워드/사용자 사전' if hit else ''
                    if not hit:
                        hit=auto_trie.longest(ts,i,source)
                        if hit:
                            key,j=hit
                            if ts[j-1]['pos'] not in NOUNS:
                                hit=None
                            else:
                                origin='corpus 자동 발견'
                    if hit:
                        key,j=hit
                        name=phrase_name(key)
                        lang=key[0][0] if len({x[0] for x in key})==1 else 'MIXED'
                        entry=dict(term=name,language=lang,source=origin,
                                   surface=source[ts[i]['start']:ts[j-1]['end']],
                                   start=ts[i]['start'],end=ts[j-1]['end'])
                        if origin.startswith('저자'): seed_hits[key]+=1
                        used.add(name)
                    else:
                        j=i+1; t=ts[i]
                        if not standalone(t):
                            i=j; continue
                        origin='단일어'
                        if t['lang']=='KO' and len(t['parts'])>1:
                            origin='원문 붙임 명사 복원'
                            orthographic[word_key(t)]+=1
                            orthographic_parts[word_key(t)]=t['parts']
                            used.add(t['lemma'])
                        entry=dict(term=t['lemma'],language=t['lang'],source=origin,
                                   surface=t['text'],start=t['start'],end=t['end'])
                    final.append(entry)
                    term=(entry['language'],entry['term'])
                    frequencies[term][field]+=1
                    doc_terms.add(term)
                    used_sources[term].add(origin)
                    i=j
                field_dfs[field].update({(e['language'], e['term']) for e in final})
                data['final_terms']=final
                counts[field]=len(final)
                # Same final terms used by the counters: no substring rematching.
                for entry in final:
                    key=(entry['language'],entry['term'])
                    h=document_hits.setdefault(key,[0,[]])
                    h[0] |= 1 if field=='title' else 2
                    if entry['surface'] not in h[1] and len(h[1])<3:
                        h[1].append(entry['surface'][:160])
            for key in keyword_by_row.get(rec['excel_row'],set()):
                h=document_hits.setdefault(key,[0,[]])
                h[0] |= 4
            dfs.update(doc_terms)
            union_dfs.update(doc_terms | keyword_by_row.get(rec['excel_row'],set()))
            paper_hits.add_document(rec['excel_row'],document_hits)
        paper_hits.commit()
        for term in keyword_df:
            frequencies[term]  # 저자키워드에만 있는 용어도 원표에 포함
            used_sources[term].add('저자키워드 등장')
        rows=[]
        for (lang,term),c in sorted(frequencies.items(),key=lambda x:(-sum(x[1].values()),x[0])):
            if sum(c.values())>=MIN_WORD_FREQ or keyword_df.get((lang,term),0)>0:
                rows.append([len(rows)+1,term,lang,c['title'],c['abstract'],sum(c.values()),
                    dfs[(lang,term)],dfs[(lang,term)]/len(records),'; '.join(sorted(used_sources[(lang,term)]))])
        cr=[]
        for c in sorted(candidates,key=lambda x:-x['freq']):
            key=c['key']; term=phrase_name(key)
            cr.append([term,' + '.join(x[1] for x in key),key[0][0],c['freq'],c['df'],c['pmi'],
                       'corpus 자동 발견'+(' + 저자키워드' if key in seeds else ''),
                       '채택' if c['eligible'] or key in seeds else '보류',
                       ' / '.join(sorted(labels.get(key,[])))])
        candidate_keys = {c['key'] for c in candidates}
        for key in sorted(seeds):
            if key not in candidate_keys:
                cr.append([phrase_name(key),' + '.join(x[1] for x in key),
                    key[0][0] if len({x[0] for x in key})==1 else 'MIXED',
                    None,None,None,'저자키워드/사용자 사전','보호', ' / '.join(sorted(labels[key]))])
        for key,c in orthographic.items():
            cr.append([key[1],' + '.join(orthographic_parts[key]),key[0],None,None,None,
                       '원문 붙임 명사 복원','보호',''])
        for r in cr:
            counter = frequencies.get((r[2], r[0]), {})
            r.extend([counter.get('title', 0), counter.get('abstract', 0), sum(counter.values())])
        config={k:v for k,v in globals().items() if k.isupper() and isinstance(v,(str,int,float,bool,list,set))}
        config={k:sorted(v) if isinstance(v,set) else v for k,v in config.items()}
        metadata=dict(input=str(Path(input_path).resolve()),mode=mode,documents=len(records),
            original_columns=header,config=config,seconds=round(time.perf_counter()-started,2),
            cache_hits=cache_hits,failed_fields=field_errors,
            error_count=len(errors),seed_count=len(seeds),auto_accepted=len(auto),
            complete=(not errors),spacy_model_version=model_version)
        fheaders=['순위','단어_용어','언어','제목빈도','초록빈도','전체빈도','문서빈도_DF','문서비율','출처']
        tables=[('전체빈도',fheaders,rows),
            ('복합어후보',['표현','구성','언어','후보창빈도','후보창_DF','최소분할_PMI','출처','상태','원키워드','최종제목빈도','최종초록빈도','최종전체빈도'],cr),
            ('오류',['Excel행','필드','키워드','오류'],[[e['excel_row'],e['field'],e['text'],e['error']] for e in errors]),
            ('실행정보',['항목','값'],[[k,json.dumps(v,ensure_ascii=False)] for k,v in metadata.items()])]
        tables.append(('저자키워드빈도',['표현','언어','저자키워드원빈도','저자키워드빈도_논문내중복제거','통합DF','제목DF','초록DF'],
            [[r[1],r[2],keyword_raw.get((r[2],r[1]),0),keyword_df.get((r[2],r[1]),0),
              union_dfs[(r[2],r[1])],field_dfs['title'][(r[2],r[1])],
              field_dfs['abstract'][(r[2],r[1])]] for r in rows]))
        # These structures are no longer needed by the report.
        del candidates, auto, seeds, labels, candidate_keys, keyword_by_row
        del seed_trie, auto_trie, frequencies, field_dfs, union_dfs, dfs
        del used_sources, orthographic, orthographic_parts, seed_hits
        del rec, data, ts, final, document_hits
        gc.collect()
        write_df_report(tables,out,web_context=dict(records=records,hits=paper_hits,errors=errors))
        print(f'완료: {out.resolve()}/analysis.xlsx')
        print(f'소요 {time.perf_counter()-started:.1f}초, 캐시 재사용 {cache_hits}개, 오류 {len(errors)}건')
        if errors:
            print('주의: 일부 필드/키워드 실패. 아래 오류를 확인하고 다시 실행하세요.')
            for error in errors: print(error)
        return out
    finally:
        if analyzer is not None:
            analyzer.cache.close()
            analyzer.ko_pos.cache_clear()
        if paper_hits is not None:
            paper_hits.close()


# ── 저자키워드 반영 및 키프레이즈 필터 ──
TITLE_WEIGHT = 5.0
ABSTRACT_WEIGHT = 1.0
KEYWORD_WEIGHT = 5.0
RANKING_VERSION = 'citation'
EXCLUDE_GENERIC_PHRASES = {
    'in vivo','in vitro','ex vivo','in silico',
    'in vivo study','in vitro study','in vivo experiment','in vitro experiment',
    'valuable insight','crucial role','important role','significant role',
    'promising approach','promising strategy','future research','further research',
    'further study','future study','present study','current study',
    'retrospective study','prospective study','retrospective cohort','prospective cohort',
    'clinical outcome','adverse event','high risk',
    'confidence interval','odd ratio','odds ratio','hazard ratio',
    'logistic regression','statistical analysis','statistical significance',
    'clinical trial','randomized controlled trial','longitudinal study',
    'primary outcome','secondary outcome','study population','human patient',
    'case report','systematic review','risk factor',
}
# 단어 경계를 지킨 접두어 필터. in vitro fertilization 등은 예외로 보존.
FILTER_GENERIC_PREFIXES = False
GENERIC_PREFIXES = {'in vivo','in vitro','ex vivo','in silico'}
KEEP_EXACT_PHRASES = {'in vitro fertilization'}


@lru_cache(maxsize=4096)
def normalized_filter_text(term):
    s=normalize_term(term).replace('‐','-').replace('‑','-')
    return re.sub(r'\s+',' ',s.replace('-',' ')).strip()


def exclusion_reason(term):
    text=normalized_filter_text(term)
    keep={normalized_filter_text(x) for x in KEEP_EXACT_PHRASES}
    if text in keep:return ''
    if text in {normalized_filter_text(x) for x in EXCLUDE_EXACT_TERMS}:return '단독 일반어'
    if text in {normalized_filter_text(x) for x in EXCLUDE_GENERIC_PHRASES}:return '일반 연구·통계·서술 표현'
    if FILTER_GENERIC_PREFIXES:
        for prefix in GENERIC_PREFIXES:
            prefix=normalized_filter_text(prefix)
            if text.startswith(prefix+' '):return f'일반 실험 조건 접두어: {prefix}'
    return ''


def lexical_word_count(term):
    return sum(bool(re.search('[A-Za-z가-힣]',w)) for w in
               re.findall(r"[A-Za-z가-힣0-9]+(?:[-'][A-Za-z가-힣0-9]+)*",str(term)))


def keyword_counters(records,labels):
    """저자키워드는 전체 항목 일치로 센다. 하위 문자열을 다시 세지 않는다."""
    alias={}
    for key,raws in labels.items():
        lang=key[0][0] if len({x[0] for x in key})==1 else 'MIXED'
        term=(lang,phrase_name(key))
        for raw in raws:alias[raw]=term
    raw_count=collections.Counter();doc_count=collections.Counter();per_row={}
    for rec in records:
        terms=[]
        for raw in rec['keywords'].split(KEYWORD_SEPARATOR):
            term=alias.get(clean(raw))
            if term is not None:terms.append(term)
        raw_count.update(terms)
        distinct=set(terms);doc_count.update(distinct)
        per_row[rec['excel_row']]=distinct
    return raw_count,doc_count,per_row


def get_document_count(tables):
    for name,headers,rows in tables:
        if name=='실행정보':
            for r in rows:
                if r[0]=='documents':
                    n=json.loads(r[1]) if isinstance(r[1],str) else r[1]
                    if isinstance(n,(int,float)) and n>=1 and int(n)==n:return int(n)
    raise ValueError('실행정보의 분석 문서 수가 필요합니다.')


def build_author_ranking(rows,stats,n,citation_stats):
    weights=[TITLE_WEIGHT,ABSTRACT_WEIGHT,KEYWORD_WEIGHT]
    if any(not math.isfinite(w) or w<0 for w in weights) or not any(weights):
        raise ValueError('가중치는 0 이상의 유한수이며 하나 이상 양수여야 합니다.')
    if WORDCLOUD_MIN_DF<1 or WORDCLOUD_MAX_WORDS<1:raise ValueError('최소 DF/표시 수는 1 이상이어야 합니다.')
    excluded=[];items=[];seen=set()
    for row in rows:
        _,term,lang,title,abstract,tf,df,ratio,source=row[:9]
        key=(lang,term)
        if key in seen:raise ValueError('중복 용어 행: '+str(key))
        seen.add(key)
        if key not in stats:raise ValueError('저자키워드 통계 누락: '+str(key))
        raw,kdf,union,title_df,abstract_df=stats[key]
        nums=[title,abstract,tf,df,raw,kdf,union,title_df,abstract_df]
        if any(not isinstance(x,(int,float)) or not math.isfinite(x) or int(x)!=x or x<0 for x in nums):
            raise ValueError('빈도는 0 이상의 정수여야 합니다: '+term)
        if title+abstract!=tf or df>tf or not max(df,kdf)<=union<=min(n,df+kdf) or union<1 or raw<kdf:
            raise ValueError('본문/키워드/통합 DF 불일치: '+term)
        if not (title_df<=title and abstract_df<=abstract and max(title_df,abstract_df)<=df<=min(n,title_df+abstract_df)):
            raise ValueError('제목·초록 DF 불일치: '+term)
        reason=exclusion_reason(term)
        if reason:
            excluded.append(list(row[:9])+[raw,kdf,union,reason]);continue
        base=TITLE_WEIGHT*title_df+ABSTRACT_WEIGHT*abstract_df+KEYWORD_WEIGHT*kdf
        idf=math.log((n+1)/(union+1))+1
        cs=citation_stats.get(key)
        if cs is None or not math.isclose(cs['base_check'],base,rel_tol=1e-12,abs_tol=1e-12):
            raise ValueError('논문별 기여점수와 DF 점수가 일치하지 않습니다: '+term)
        if not math.isfinite(cs['citation_score']):
            raise ValueError('인용 가중 점수가 너무 큽니다. CITATION_ALPHA를 줄이세요.')
        words=lexical_word_count(term)
        items.append(dict(term=term,lang=lang,title=title,abstract=abstract,tf=tf,df=df,
            raw=raw,kdf=kdf,union=union,title_df=title_df,abstract_df=abstract_df,ratio=union/n,source=source,words=words,phrase=words>=2,
            base=base,idf=idf,with_idf=base*idf,**cs,citation_with_idf=cs['citation_score']*idf))
    orders={};frequencies={}
    for version,score in SCORE_FIELDS.items():
        ordered=sorted(items,key=lambda x:(not x['phrase'],-x[score],-x['union'],-x['tf'],x['lang'],x['term']))
        eligible=[x for x in ordered if x['phrase'] and x['union']>=WORDCLOUD_MIN_DF and x[score]>0][:WORDCLOUD_MAX_WORDS]
        chosen={(x['lang'],x['term']) for x in eligible}
        labels=collections.Counter(x['term'] for x in eligible)
        maximum=max((x[score] for x in eligible),default=1.0)
        f={}
        for rank,x in enumerate(ordered,1):
            x['rank_'+version]=rank
            selected=(x['lang'],x['term']) in chosen
            x['selected_'+version]=selected
            label=x['term']+(f" ({x['lang']})" if labels[x['term']]>1 else '')
            x['label_'+version]=label
            x['relative_'+version]=x[score]/maximum if selected else None
            if selected:f[label]=x[score]
        orders[version]=ordered;frequencies[version]=f
    return orders,excluded,frequencies


def write_df_report(tables,out,web_context=None):
    out=Path(out);out.mkdir(parents=True,exist_ok=True)
    if web_context is None:
        raise ValueError('인용 점수에는 논문별 연결이 필요합니다. 원본 Excel을 분석하세요.')
    if RANKING_VERSION not in SCORE_FIELDS:
        raise ValueError('RANKING_VERSION은 no_idf, idf, citation, citation_idf 중 하나입니다.')
    n=get_document_count(tables)
    kwtable=next(t for t in tables if t[0]=='저자키워드빈도')
    needed=['표현','언어','저자키워드원빈도','저자키워드빈도_논문내중복제거','통합DF','제목DF','초록DF']
    indexes=[kwtable[1].index(k) for k in needed]
    stats={}
    for row in kwtable[2]:
        r=[row[i] for i in indexes];stats[(r[1],r[0])]=tuple(r[2:])
    rows=next(t[2] for t in tables if t[0]=='전체빈도')
    orders,excluded,frequencies=build_author_ranking(rows,stats,n,citation_term_stats(web_context))
    statuses={};web_clouds={}
    for version in SCORE_FIELDS:
        target=out/f'wordcloud_{version}.png'
        if target.exists():target.unlink()
        try:
            rendered,layout=render_df_wordcloud(frequencies[version],target)
            web_clouds[version]=layout.pop('_web')
            statuses[version]=f'완료: {len(rendered)}/{len(frequencies[version])}개 표시'
            if layout['표시되지않은각도']:
                statuses[version]+=f" / 미표시 각도 {layout['표시되지않은각도']}"
        except Exception as e:
            if target.exists():target.unlink()
            statuses[version]=f'PNG 생성 실패: {type(e).__name__}: {e}'
            print(statuses[version])
    columns=['순위','표현','제목빈도','제목DF','초록빈도','초록DF','전체빈도합계','저자키워드빈도','저자키워드DF']
    final=[]
    names={'no_idf':'IDF미적용','idf':'IDF적용','citation':'인용반영_IDF미적용','citation_idf':'인용반영_IDF적용'}
    for version,field in SCORE_FIELDS.items():
        use_idf=version in {'idf','citation_idf'}
        values=[]
        for x in orders[version]:
            values.append([x['rank_'+version],x['term'],x['title'],x['title_df'],x['abstract'],x['abstract_df'],
                           x['tf']+x['raw'],x['raw'],x['kdf']]+([x['idf']] if use_idf else [])+[x[field]])
        final.append((names[version],columns+(['IDF'] if use_idf else [])+['순위점수'],values))
    baseline=orders['no_idf']
    final.append(('두버전비교',['표현','제목DF','초록DF','저자키워드DF','IDF','순위점수',
                     'IDF 적용 순위점수','미적용순위','적용순위','순위상승폭'],
        [[x['term'],x['title_df'],x['abstract_df'],x['kdf'],x['idf'],x['base'],x['with_idf'],
          x['rank_no_idf'],x['rank_idf'],x['rank_no_idf']-x['rank_idf']] for x in baseline]))
    final.append(('인용비교',['표현','통합DF','피인용합계','피인용논문수','피인용누락논문수',
                     '인용미반영점수','인용반영점수','인용미반영순위','인용반영순위','인용순위상승폭',
                     'IDF','IDF_인용미반영점수','IDF_인용반영점수','IDF_인용미반영순위','IDF_인용반영순위','IDF_인용순위상승폭'],
        [[x['term'],x['union'],x['citation_total'],x['cited_papers'],x['citation_missing'],
          x['base'],x['citation_score'],x['rank_no_idf'],x['rank_citation'],x['rank_no_idf']-x['rank_citation'],
          x['idf'],x['with_idf'],x['citation_with_idf'],x['rank_idf'],x['rank_citation_idf'],
          x['rank_idf']-x['rank_citation_idf']] for x in baseline]))
    final.append(('단일어참고',['순위','표현','구분','출처'],
                  [[x['rank_'+RANKING_VERSION],x['term'],'단일어',x['source']]
                   for x in orders[RANKING_VERSION] if not x['phrase']]))
    final.append(('제외표현',['원순위','표현','출처','사유'],[[r[0],r[1],r[8],r[-1]] for r in excluded]))
    for name,head,data in tables:
        if name not in {'전체빈도','복합어후보'}:continue
        wanted={'전체빈도':['순위','표현','제목빈도','초록빈도','문서빈도(DF합계)','출처'],
                '복합어후보':['표현','구성','출처','상태','원키워드']}[name]
        # Preserve previous export: this DF is title/abstract UNION, not their DF sum.
        alias={'표현':'단어_용어','문서빈도(DF합계)':'문서빈도_DF'} if name=='전체빈도' else {}
        indices=[head.index(alias.get(h,h)) for h in wanted]
        final.append((name,wanted,[[r[i] for i in indices] for r in data]))
    write_excel(out/'analysis.xlsx',final)
    write_explorer(out,orders,web_clouds,web_context)
    print(f'기본점수 = {TITLE_WEIGHT:g}×제목DF + {ABSTRACT_WEIGHT:g}×초록DF + {KEYWORD_WEIGHT:g}×저자키워드DF')
    print(f'인용점수 = Σ[논문별 기여점수 × (1 + {CITATION_ALPHA:g} × ln(1+피인용횟수))]')
    print('이미지 생성 상태:',statuses)
    print(f'완료: 논문 {n}편, 제외 {len(excluded)}개, 키프레이즈 {sum(x["phrase"] for x in baseline)}개')
    return out


# ── 대각선 워드클라우드 설정: 3번 셀에서 변경 가능 ──
WORDCLOUD_ANGLES = [0, 30, 45, 90]
WORDCLOUD_ANGLE_PROBS = [0.45, 0.20, 0.20, 0.15]
WORDCLOUD_GAP = 5
WORDCLOUD_MAX_FONT = 150
WORDCLOUD_MIN_FONT = 16
WORDCLOUD_SIZE_POWER = 1.0
WORDCLOUD_ATTEMPTS = 1600
import random
def load_render_libraries():
    global np, pd, Image, ImageDraw, ImageFont, ImageFilter, matplotlib
    import numpy as np
    import pandas as pd
    from PIL import Image, ImageDraw, ImageFont, ImageFilter
    import matplotlib
from collections import Counter
from tqdm.auto import tqdm

def rotated_mask(text, font_path, size, angle, gap):
    font = ImageFont.truetype(font_path, size)
    box = font.getbbox(text)
    w, h = box[2]-box[0], box[3]-box[1]
    if w <= 0 or h <= 0: raise ValueError(f'표현을 그릴 수 없습니다: {text!r}')
    mask = Image.new('L', (w+8, h+8), 0)
    ImageDraw.Draw(mask).text((4-box[0],4-box[1]), text, font=font, fill=255)
    mask = mask.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True)
    bbox = mask.getbbox()
    if bbox is None: raise ValueError(f'빈 글자 이미지: {text!r}')
    mask = mask.crop(bbox)
    # 최종 표시 픽셀과 같은 마스크를 검사. 작은 안티앨리어싱 픽셀도 포함.
    padded = Image.new('L', (mask.width+2*gap, mask.height+2*gap), 0)
    padded.paste(mask, (gap,gap))
    ink = np.asarray(padded) > 0
    guard = Image.fromarray((ink.astype(np.uint8)*255))
    if gap: guard = guard.filter(ImageFilter.MaxFilter(2*gap+1))
    return padded, ink, np.asarray(guard)>0


def pack_rows(mask):
    return [int.from_bytes(row.tobytes(),'little') for row in
            np.packbits(mask, axis=1, bitorder='little')]


def scheduled_angles(n, angles, probs, seed):
    probs = np.asarray(probs, dtype=float)
    if len(angles)!=len(probs) or not len(angles) or len(set(angles))!=len(angles):
        raise ValueError('ANGLES는 중복 없이, ANGLE_PROBS와 같은 길이로 지정하세요.')
    if not np.all(np.isfinite(probs)) or np.any(probs<=0):
        raise ValueError('각 각도의 비중은 양수여야 합니다.')
    if not all(math.isfinite(a) for a in angles): raise ValueError('각도는 유한수여야 합니다.')
    # 충분한 표현이 있으면 모든 각도에 적어도 한 개씩 배정.
    rng = random.Random(seed)
    result = list(angles[:min(n,len(angles))])
    result += rng.choices(angles, weights=probs.tolist(), k=n-len(result))
    rng.shuffle(result)
    return result


def render_diagonal(df, output_png, *, font_path, width=2000, height=1200,
                    angles=(0,30,45,90), angle_probs=(.45,.2,.2,.15),
                    colormap='hsv', background='white', gap=5,
                    max_font=150, min_font=16, size_power=1.0, seed=42,
                    attempts=1600, verify=True):
    load_render_libraries()
    if width<100 or height<100 or not 1<=min_font<=max_font or gap<0 or attempts<1:
        raise ValueError('캔버스·글꼴·여백·시도 횟수 설정을 확인하세요.')
    if not math.isfinite(size_power) or size_power<=0: raise ValueError('SIZE_POWER는 양수여야 합니다.')
    cmap = matplotlib.colormaps[colormap]
    canvas = Image.new('RGB',(width,height),background)
    hit_canvas = Image.new('RGB',(width,height),(0,0,0))
    occupied = [0]*height
    # 별도 최종 검증용 배열. 단어별 최종 픽셀 마스크를 다시 누적해 검사.
    coverage = np.zeros((height,width),dtype=bool) if verify else None
    rng = random.Random(seed)
    assignments = scheduled_angles(len(df),angles,angle_probs,seed)
    max_weight = float(df['weight'].max())
    previous_size = max_font
    records=[]; started=time.perf_counter()
    for (rank, row), angle in tqdm(zip(enumerate(df.to_dict('records'),1),assignments),
                                  total=len(df),desc=Path(output_png).stem):
        term = row['label']; weight=float(row['weight'])
        target = max(min_font, min(max_font, round(max_font*(weight/max_weight)**size_power)))
        # 점수순 글자 크기의 역전을 방지. 공간 부족으로 축소한 크기는 기록.
        size = min(target,previous_size)
        placed=False
        while size>=min_font:
            mask, ink, guard = rotated_mask(term,font_path,size,angle,gap)
            h,w = ink.shape
            if w<=width and h<=height:
                packed=pack_rows(ink); guard_rows=pack_rows(guard)
                row_order=sorted([i for i,r in enumerate(packed) if r],
                                 key=lambda i:packed[i].bit_count(),reverse=True)
                maxx,maxy=width-w,height-h
                for attempt in range(attempts):
                    if attempt==0:
                        x,y=maxx//2,maxy//2
                    elif attempt < attempts//2:
                        # 중심 부근부터 탐색 범위를 넓힙니다.
                        fraction=(attempt/(attempts//2))**.5
                        x=round(maxx/2+rng.uniform(-1,1)*maxx/2*fraction)
                        y=round(maxy/2+rng.uniform(-1,1)*maxy/2*fraction)
                    else:
                        x,y=rng.randint(0,maxx),rng.randint(0,maxy)
                    if any(occupied[y+i] & (packed[i]<<x) for i in row_order): continue
                    if verify:
                        roi=coverage[y:y+h,x:x+w]
                        if np.any(roi & ink): raise AssertionError('최종 픽셀 겹침 검증 실패')
                        roi |= ink
                    for i,bits in enumerate(guard_rows): occupied[y+i] |= bits<<x
                    # 표현별 색 고정: 두 IDF 버전에서 같은 표현은 같은 색.
                    color_seed=int.from_bytes(hashlib.sha256(f'{seed}:{term}'.encode()).digest()[:8],'big')
                    rgba=cmap(random.Random(color_seed).random())
                    color=tuple(round(v*255) for v in rgba[:3])
                    canvas.paste(color,(x,y,x+w,y+h),mask)
                    # Exact ink pixels, not overlapping rectangular click boxes.
                    hit_canvas.paste((rank & 255,(rank>>8)&255,(rank>>16)&255),
                                     (x,y,x+w,y+h),Image.fromarray(ink.astype(np.uint8)*255))
                    placed=True;previous_size=size
                    break
            if placed: break
            if size==min_font:break
            size=max(min_font,size-max(2,round(size*.12)))
        records.append({'표현':row['표현'],'언어':row['언어'],'표시문구':term,
                        '선정순서':rank,'입력점수':weight,'각도':angle,
                        '목표글꼴크기':target,'실제글꼴크기':size if placed else None,
                        'PNG실제표시':placed,'x':x if placed else None,'y':y if placed else None,
                        '상태':'완료' if placed else '최소 글꼴 크기에서도 배치 공간 부족'})
    canvas.save(output_png)
    placed_angles=Counter(r['각도'] for r in records if r['PNG실제표시'])
    summary={'입력표현수':len(records),'표시표현수':sum(placed_angles.values()),
             '각도별표시수':dict(placed_angles),'표시되지않은각도':[a for a in angles if not placed_angles[a]],
             '겹침검증':'통과' if verify else '미실행','소요초':round(time.perf_counter()-started,2)}
    buf=io.BytesIO();hit_canvas.save(buf,format='PNG')
    summary['_web']=dict(hit='data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode('ascii'),
                         labels={str(r['선정순서']):r['표시문구'] for r in records if r['PNG실제표시']})
    return pd.DataFrame(records),summary



def render_df_wordcloud(frequencies, output_file):
    load_render_libraries()
    output_file=Path(output_file)
    if not frequencies:
        raise ValueError('표시할 키프레이즈가 없습니다. 최소DF·제외 목록을 확인하세요.')
    data=[dict(표현=term, 언어='', label=term, weight=score)
          for term,score in sorted(frequencies.items(),key=lambda kv:(-kv[1],kv[0]))]
    df=pd.DataFrame(data)
    placement,summary=render_diagonal(
        df,output_file,font_path=choose_wordcloud_font(frequencies),
        width=WORDCLOUD_WIDTH,height=WORDCLOUD_HEIGHT,
        angles=WORDCLOUD_ANGLES,angle_probs=WORDCLOUD_ANGLE_PROBS,
        colormap=WORDCLOUD_COLORMAP,background=WORDCLOUD_BACKGROUND,gap=WORDCLOUD_GAP,
        max_font=WORDCLOUD_MAX_FONT,min_font=WORDCLOUD_MIN_FONT,
        size_power=WORDCLOUD_SIZE_POWER,seed=WORDCLOUD_RANDOM_STATE,
        attempts=WORDCLOUD_ATTEMPTS,verify=True)
    print(output_file.name, {k:v for k,v in summary.items() if k!='_web'})
    return set(placement.loc[placement['PNG실제표시'],'표시문구']),summary

EXPLORER_TEMPLATE = re.search(r'<!--BEGIN_REPORT-->\n(.*?)\n<!--END_REPORT-->', (Path(__file__).with_name('ui')/'index.html').read_text(encoding='utf-8'), re.S).group(1)

def write_explorer(out,orders,clouds,context):
    records=context['records']; hits=context['hits']
    row_index={r['excel_row']:i for i,r in enumerate(records)}
    def plain(value):
        return html.unescape(re.sub(r'<[^>]*>','',str(value or ''))).strip()
    papers=[]
    for r in records:
        doi=re.sub(r'^https?://(?:dx\.)?doi\.org/|^doi:\s*','',r['doi'].strip(),flags=re.I)
        link=r['link'].strip()
        if urlparse(link).scheme.lower() not in {'http','https'}: link=''
        papers.append(dict(row=r['excel_row'],id=plain(r['id']),title=plain(r['title']),
            authors=plain(r['authors']),year=plain(r['year']),journal=plain(r['journal']),
            doi=doi,link=link,source=plain(r['source']),keywords=plain(r['keywords']),
            abstract=clean(r['abstract']) if HTML_INCLUDE_ABSTRACT else '',citations=r['citations']))
    terms=[]; labels={v:{} for v in SCORE_FIELDS}
    for x in orders['no_idf']:
        if not any(x['selected_'+v] for v in SCORE_FIELDS):continue
        key=(x['lang'],x['term']); matches=hits.get(key,{})
        # Fail visibly if the website and analysis would report different papers.
        actual=(len(matches),sum(bool(h[0]&1) for h in matches.values()),
                sum(bool(h[0]&2) for h in matches.values()),sum(bool(h[0]&4) for h in matches.values()))
        expected=(x['union'],x['title_df'],x['abstract_df'],x['kdf'])
        if actual!=expected:raise AssertionError(f'논문 연결/DF 불일치: {key}: {actual} != {expected}')
        term_id=str(len(terms))
        t=dict(key=term_id,term=x['term'],language=x['lang'],union=x['union'],
               titleDF=x['title_df'],abstractDF=x['abstract_df'],keywordDF=x['kdf'],idf=x['idf'],
               score={v:x[field] for v,field in SCORE_FIELDS.items()},
               rank={v:x['rank_'+v] for v in labels},selected={v:x['selected_'+v] for v in labels},
               hits=[[row_index[row],h[0],' / '.join(h[1])] for row,h in sorted(matches.items())])
        terms.append(t)
        for v in labels:
            if x['selected_'+v]:labels[v][x['label_'+v]]=term_id
    for v,cloud in clouds.items():
        cloud['keys']={rank:labels[v][label] for rank,label in cloud.pop('labels').items()}
        cloud['png']='data:image/png;base64,'+base64.b64encode((out/f'wordcloud_{v}.png').read_bytes()).decode('ascii')
        cloud['placed']=len(cloud['keys'])
    warnings=[]
    missing=sum(r['citations'] is None for r in records)
    if missing: warnings.append(f'피인용 횟수 누락 {missing}건: 인용 가중치 1 적용. 실제 피인용 0과 구분됩니다.')
    if context['errors']:
        examples='; '.join(f"Excel {e['excel_row']}행 {e['field']}: {e['error']}" for e in context['errors'][:5])
        warnings.append(f"분석 오류 {len(context['errors'])}건이 있어 일부 내용이 집계되지 않았습니다. {examples}")
    identity=[r['id'].strip().casefold() for r in records]
    duplicates=len(identity)-len(set(identity))
    if duplicates:warnings.append(f'동일 식별자가 중복된 행이 {duplicates}개 있습니다. 기존 점수와 맞추기 위해 Excel 행 단위를 유지했습니다. 원본 중복을 정리한 뒤 다시 분석하세요.')
    if len(clouds)<len(SCORE_FIELDS):warnings.append('일부 워드클라우드 이미지가 생성되지 않았습니다. 표현 선택 메뉴는 사용할 수 있습니다.')
    payload=dict(citationAlpha=CITATION_ALPHA,citationAsOf=CITATION_AS_OF,defaultVersion=RANKING_VERSION,papers=papers,terms=terms,clouds=clouds,width=WORDCLOUD_WIDTH,height=WORDCLOUD_HEIGHT,
                 source=', '.join(sorted({p['source'] for p in papers if p['source']})) or '업로드 파일',
                 weights=[TITLE_WEIGHT,ABSTRACT_WEIGHT,KEYWORD_WEIGHT],created=datetime.now().strftime('%Y-%m-%d'),
                 warnings=warnings)
    # Escape HTML-sensitive characters so even an abstract containing </script> stays data.
    before, separator, after = EXPLORER_TEMPLATE.partition('__PAYLOAD__')
    if not separator:
        raise ValueError('ui/index.html의 보고서 템플릿에 __PAYLOAD__가 없습니다.')
    with (out/'paper_explorer.html').open('w', encoding='utf-8') as stream:
        stream.write(before)
        for chunk in json.JSONEncoder(ensure_ascii=False, allow_nan=False).iterencode(payload):
            stream.write(chunk.replace('&',r'\u0026').replace('<',r'\u003c').replace('>',r'\u003e'))
        stream.write(after)
    print(f'클릭형 HTML 완료: {len(terms)}개 표현, {len(papers)}건. DF/논문 연결 검증 통과.')

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input',help='논문 Excel 파일 경로')
    parser.add_argument('--mode',choices=['sample','full','existing'],default='sample')
    parser.add_argument('--output-dir',default='nlp_output_df')
    parser.add_argument('--cache-dir',default='nlp_cache')
    args=parser.parse_args()
    if args.mode == 'existing':
        reprocess_existing(args.input,args.output_dir)
    else:
        run(args.input,args.mode,args.output_dir,args.cache_dir)