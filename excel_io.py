"""Read-only workbook inspection before expensive NLP."""
from io import BytesIO
from zipfile import ZipFile,BadZipFile
from collections import Counter
import openpyxl

MAX_ROWS=50000

def workbook_sheets(content):
    if len(content)>50*1024*1024:raise ValueError('파일 크기는 50MB 이하여야 합니다.')
    try:
        with ZipFile(BytesIO(content)) as z:
            if 'xl/workbook.xml' not in z.namelist():raise ValueError('올바른 Excel(.xlsx) 파일이 아닙니다.')
            if sum(f.file_size for f in z.infolist())>400*1024*1024:raise ValueError('압축 해제한 Excel 크기가 너무 큽니다. 분석할 시트만 새 파일로 저장해 주세요.')
        wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
        try:return wb.sheetnames
        finally:wb.close()
    except (BadZipFile,KeyError) as e:raise ValueError('Excel 파일을 읽을 수 없습니다. .xlsx 형식으로 다시 저장해 주세요.') from e

def inspect_sheet(content,sheet):
    wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
    try:
        rows=wb[sheet].iter_rows(values_only=True)
        first=next(rows,None)
        if first is None:raise ValueError('선택한 시트가 비어 있습니다.')
        header=[str(v).strip() if v is not None else '' for v in first]
        duplicate=[k for k,n in Counter(header).items() if k and n>1]
        if duplicate:raise ValueError('중복된 열 이름이 있습니다: '+', '.join(duplicate))
        if not any(header):raise ValueError('첫 번째 행에 열 이름이 필요합니다.')
        count=0;preview=[]
        for index,row in enumerate(rows,2):
            if index>MAX_ROWS+1:raise ValueError(f'한 시트는 최대 {MAX_ROWS:,}행까지 지원합니다. 불필요한 빈 행을 제거해 주세요.')
            if not any(v is not None for v in row):continue
            count+=1
            if len(preview)<5:preview.append({h:v for h,v in zip(header,row) if h})
        if not count:raise ValueError('분석할 논문 행이 없습니다.')
        return {'columns':[h for h in header if h],'count':count,'preview':preview}
    finally:wb.close()
